"""Generate Nova Candles demo workbooks and safely sync products/partners to Odoo."""

from __future__ import annotations

import argparse
import hashlib
import math
import os
import random
import sys
import xmlrpc.client
from datetime import date, datetime, timedelta
from pathlib import Path
from typing import Any

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.table import Table, TableStyleInfo


SEED_DEFAULT = 20260927
BATCH_SIZE = 100
SHEET_PRODUCTS = "Products"
SHEET_PARTNERS = "Partners"

SHEET_ALIASES = {
    SHEET_PRODUCTS: ("Products", "Product", "المنتجات", "المنتجات الرئيسية", "قاعدة المنتجات"),
    SHEET_PARTNERS: (
        "Partners", "Customers", "Distributors", "العملاء", "العملاء والموزعين",
        "الموزعين والعملاء", "الموزعين", "نقاط البيع",
    ),
}

PRODUCT_HEADER_ALIASES = {
    "SKU": ("SKU", "Internal Reference", "Default Code", "كود المنتج", "كود الصنف", "الرمز الداخلي"),
    "Name": ("Name", "Product Name", "اسم المنتج", "الاسم"),
    "Cost EGP": ("Cost EGP", "Cost", "Standard Price", "التكلفة", "سعر التكلفة", "سعر التكلفة جنيه"),
    "Retail Price EGP": (
        "Retail Price EGP", "Sale Price", "Sales Price", "Retail Price", "سعر البيع",
        "سعر البيع جنيه", "سعر البيع ج م", "السعر",
    ),
    "Quantity On Hand": (
        "Quantity On Hand", "Stock", "Inventory", "المخزون", "الكمية", "رصيد المخزون",
    ),
    "Product Category": ("Product Category", "Category", "الفئة", "التصنيف"),
    "Product Type": ("Product Type", "Type", "نوع المنتج", "الخامة"),
    "Wick Type": ("Wick Type", "نوع الفتلة", "الفتلة"),
    "Wick Type Code": ("Wick Type Code", "Wick Code"),
    "Fragrance Oil": ("Fragrance Oil", "الزيت العطري", "العطر"),
    "Materials": ("Materials", "الخامات", "المواد"),
    "Sales Description": ("Sales Description", "Description", "الوصف", "الوصف المختصر"),
    "Wholesale Price EGP": ("Wholesale Price EGP", "Wholesale Price", "سعر الجملة"),
}

PARTNER_HEADER_ALIASES = {
    "Customer Code": ("Customer Code", "Partner Code", "Reference", "كود العميل", "كود الموزع"),
    "Name": ("Name", "Customer Name", "Distributor Name", "اسم العميل", "اسم العميل / الشركة", "اسم الموزع", "اسم المعرض"),
    "City": ("City", "المدينة", "المدينة / المحافظة", "المدينة / العنوان"),
    "Governorate": ("Governorate", "المحافظة"),
    "District": ("District", "Area", "المنطقة", "الحي"),
    "Phone": ("Phone", "Mobile", "رقم الهاتف", "الهاتف", "الموبايل"),
    "Customer Segment": ("Customer Segment", "Category", "Classification", "التصنيف", "نوع العميل"),
    "Sales Channel": ("Sales Channel", "Channel", "قناة البيع"),
    "Contact Person": ("Contact Person", "اسم المسؤول", "المسؤول"),
    "Credit Limit EGP": ("Credit Limit EGP", "Credit Limit", "الحد الائتماني"),
}

AROMA_OILS = [
    ("لافندر فرنسي", "French Lavender"),
    ("فانيليا", "Vanilla"),
    ("عنبر", "Amber"),
    ("خشب الصندل", "Sandalwood"),
    ("ياسمين", "Jasmine"),
    ("ورد دمشقي", "Damask Rose"),
    ("مسك أبيض", "White Musk"),
    ("عود ملكي", "Royal Oud"),
    ("تين بري", "Wild Fig"),
    ("قرفة وقرنفل", "Cinnamon Clove"),
    ("توت بري", "Wild Berries"),
    ("حمضيات متوسطية", "Mediterranean Citrus"),
    ("زهر البرتقال", "Orange Blossom"),
    ("شاي أخضر", "Green Tea"),
    ("نعناع", "Mint"),
    ("قهوة عربية", "Arabic Coffee"),
    ("كراميل مملح", "Salted Caramel"),
    ("جوز الهند", "Coconut"),
    ("زهرة القطن", "Cotton Flower"),
    ("بحر ونسمات", "Sea Breeze"),
]

PRODUCT_FORMS = [
    ("صويا طبيعية", "Soy"),
    ("جل شفاف", "Gel"),
    ("شمع نباتي", "Plant Wax"),
    ("شمع نحل", "Beeswax"),
]

VESSELS = [
    ("وعاء زجاجي", "Glass"),
    ("وعاء خرساني", "Concrete"),
    ("وعاء سيراميك", "Ceramic"),
    ("قالب معدني", "Metal Tin"),
]

WICKS = [("خشبية", "wood"), ("قطنية", "cotton")]

EGYPT_GOVERNORATES: dict[str, list[str]] = {
    "القاهرة": ["مدينة نصر", "المعادي", "مصر الجديدة", "التجمع الخامس", "وسط البلد"],
    "الجيزة": ["الدقي", "المهندسين", "الشيخ زايد", "6 أكتوبر", "فيصل"],
    "الإسكندرية": ["سموحة", "سيدي جابر", "ميامي", "العجمي", "المنتزه"],
    "القليوبية": ["بنها", "شبرا الخيمة", "العبور", "قليوب"],
    "الشرقية": ["الزقازيق", "العاشر من رمضان", "بلبيس", "منيا القمح"],
    "الدقهلية": ["المنصورة", "ميت غمر", "طلخا", "دكرنس"],
    "الغربية": ["طنطا", "المحلة الكبرى", "زفتى", "كفر الزيات"],
    "المنوفية": ["شبين الكوم", "منوف", "أشمون", "السادات"],
    "البحيرة": ["دمنهور", "كفر الدوار", "رشيد", "إدكو"],
    "كفر الشيخ": ["كفر الشيخ", "دسوق", "بلطيم", "فوه"],
    "دمياط": ["دمياط", "رأس البر", "فارسكور", "دمياط الجديدة"],
    "بورسعيد": ["حي الشرق", "حي العرب", "بورفؤاد", "حي المناخ"],
    "الإسماعيلية": ["الإسماعيلية", "فايد", "القنطرة", "التل الكبير"],
    "السويس": ["حي السويس", "عتاقة", "الأربعين", "فيصل"],
    "الفيوم": ["الفيوم", "سنورس", "إطسا", "طامية"],
    "بني سويف": ["بني سويف", "ناصر", "إهناسيا", "ببا"],
    "المنيا": ["المنيا", "ملوي", "سمالوط", "مغاغة"],
    "أسيوط": ["أسيوط", "ديروط", "القوصية", "أبنوب"],
    "سوهاج": ["سوهاج", "طهطا", "جرجا", "أخميم"],
    "قنا": ["قنا", "نجع حمادي", "قوص", "دشنا"],
    "الأقصر": ["الأقصر", "القرنة", "إسنا", "الطود"],
    "أسوان": ["أسوان", "كوم أمبو", "إدفو", "دراو"],
    "البحر الأحمر": ["الغردقة", "سفاجا", "مرسى علم", "القصير"],
    "الوادي الجديد": ["الخارجة", "الداخلة", "الفرافرة", "باريس"],
    "مطروح": ["مرسى مطروح", "العلمين", "الضبعة", "سيوة"],
    "شمال سيناء": ["العريش", "بئر العبد", "رفح", "الشيخ زويد"],
    "جنوب سيناء": ["شرم الشيخ", "الطور", "دهب", "نويبع"],
}

STORE_PREFIXES = [
    "بيت الشموع",
    "لمسة عطر",
    "ركن الهدايا",
    "دار الروائح",
    "معرض النور",
    "أطياف المنزل",
    "سوق الصفوة",
    "عالم الديكور",
    "مخازن الندى",
    "لمسة فاخرة",
]

FIRST_NAMES = [
    "أحمد", "محمد", "محمود", "مصطفى", "عمر", "يوسف", "إبراهيم", "خالد",
    "عمرو", "حسن", "حسين", "طارق", "ياسر", "كريم", "مروان", "سارة", "منى",
    "مريم", "نور", "هبة", "دينا", "آية", "مي", "رانيا", "سلمى", "ياسمين",
]
MIDDLE_NAMES = [
    "محمد", "محمود", "حسن", "حسين", "إبراهيم", "عبد الرحمن", "عبد الله",
    "مصطفى", "علي", "أحمد", "سعيد", "فؤاد", "جمال", "كمال", "صالح",
]
FAMILY_NAMES = [
    "المنشاوي", "الشربيني", "الرفاعي", "السيد", "عبد الرازق", "حمدي", "عوض",
    "النجار", "البدري", "سليمان", "مراد", "إسماعيل", "الحداد", "فوزي", "زكي",
]

DEPARTMENTS = [
    ("الإنتاج", ["مشغل إنتاج", "فني تصنيع شموع", "مشرف خط إنتاج", "عامل تجهيز خامات"]),
    ("الجودة", ["مفتش جودة", "أخصائي جودة", "مشرف جودة", "فني اختبارات احتراق"]),
    ("المبيعات", ["مندوب مبيعات", "مسؤول حسابات", "مشرف مبيعات", "مدير منطقة"]),
    ("أودو وتقنية المعلومات", ["مسؤول أودو", "محلل نظم", "دعم فني", "مسؤول بيانات"]),
    ("الحسابات", ["محاسب", "محاسب مخزون", "مراجع حسابات", "مدير حسابات"]),
]


def _round_to_five(value: float) -> int:
    return int(round(value / 5) * 5)


def _synthetic_name(rng: random.Random) -> str:
    return " ".join(
        [
            rng.choice(FIRST_NAMES),
            rng.choice(MIDDLE_NAMES),
            rng.choice(MIDDLE_NAMES),
            rng.choice(FAMILY_NAMES),
        ]
    )


def generate_products(rng: random.Random, count: int = 1000) -> list[dict[str, Any]]:
    products: list[dict[str, Any]] = []
    for index in range(1, count + 1):
        oil_ar, oil_en = rng.choice(AROMA_OILS)
        form_ar, form_en = rng.choice(PRODUCT_FORMS)
        vessel_ar, vessel_en = rng.choice(VESSELS)
        wick_ar, wick_code = rng.choice(WICKS)
        cost = _round_to_five(rng.uniform(45, 420))
        retail = _round_to_five(cost * rng.uniform(2.0, 2.7))
        wholesale = _round_to_five(retail * rng.uniform(0.72, 0.84))
        product_name = f"شمعة {oil_ar} - {vessel_ar} {index:04d}"
        product_category = f"Nova {form_en} Candles"
        materials = f"{form_ar}، {vessel_ar}، فتلة {wick_ar}"
        products.append(
            {
                "SKU": f"NC-ENT-{index:05d}",
                "Name": f"{product_name} ({oil_en} {vessel_en})",
                "Product Type": form_ar,
                "Wick Type": wick_ar,
                "Wick Type Code": wick_code,
                "Fragrance Oil": f"زيت {oil_ar} ({oil_en})",
                "Cost EGP": cost,
                "Wholesale Price EGP": wholesale,
                "Retail Price EGP": retail,
                "Quantity On Hand": rng.randint(0, 300),
                "Product Category": product_category,
                "Materials": materials,
                "Sales Description": f"شمعة {form_ar} معطرة بزيت {oil_ar} في {vessel_ar}.",
                "Synthetic Demo Data": "YES",
            }
        )
    return products


def generate_partners(rng: random.Random, count: int = 500) -> list[dict[str, Any]]:
    governorates = list(EGYPT_GOVERNORATES.items())
    partners: list[dict[str, Any]] = []
    for index in range(1, count + 1):
        governorate, districts = rng.choice(governorates)
        district = rng.choice(districts)
        segment = rng.choices(["جملة", "تجزئة", "أونلاين"], weights=[35, 50, 15])[0]
        store_prefix = rng.choice(STORE_PREFIXES)
        partners.append(
            {
                "Customer Code": f"NC-CUS-DEMO-{index:04d}",
                "Name": f"{store_prefix} {district} {index:03d}",
                "Customer Segment": segment,
                "Sales Channel": {"جملة": "موزع جملة", "تجزئة": "نقطة بيع", "أونلاين": "متجر إلكتروني"}[segment],
                "Governorate": governorate,
                "District": district,
                "Contact Person": _synthetic_name(rng),
                "Phone": f"SIM-+20-10-0000-{index:04d}",
                "Credit Limit EGP": _round_to_five(rng.uniform(5000, 250000)),
                "Synthetic Demo Data": "YES",
            }
        )
    return partners


def generate_sales(
    rng: random.Random,
    products: list[dict[str, Any]],
    partners: list[dict[str, Any]],
    count: int = 1000,
) -> list[dict[str, Any]]:
    first_date = date(2024, 1, 1)
    last_date = date.today()
    day_range = max((last_date - first_date).days, 1)
    sales: list[dict[str, Any]] = []
    for index in range(1, count + 1):
        partner = rng.choice(partners)
        product = rng.choice(products)
        quantity = rng.randint(1, 24)
        price_key = "Wholesale Price EGP" if partner["Customer Segment"] == "جملة" else "Retail Price EGP"
        unit_price = product[price_key]
        transaction_date = first_date + timedelta(days=rng.randint(0, day_range))
        sales.append(
            {
                "Invoice Number": f"DEMO-INV-{transaction_date.year}-{index:06d}",
                "Transaction Date": transaction_date,
                "Customer Code": partner["Customer Code"],
                "Customer Name": partner["Name"],
                "Sales Channel": partner["Sales Channel"],
                "SKU": product["SKU"],
                "Product Name": product["Name"],
                "Quantity": quantity,
                "Unit Price EGP": unit_price,
                "Total EGP": quantity * unit_price,
                "Payment Status": rng.choices(
                    ["مدفوع", "مدفوع جزئيًا", "معلق"], weights=[65, 20, 15]
                )[0],
                "Synthetic Demo Data": "YES",
            }
        )
    return sales


def generate_employees(rng: random.Random, count: int = 300) -> list[dict[str, Any]]:
    today = date.today()
    employees: list[dict[str, Any]] = []
    for index in range(1, count + 1):
        department, titles = rng.choice(DEPARTMENTS)
        if department in {"الإنتاج", "الجودة"}:
            salary = rng.randint(7000, 18000)
        elif department == "المبيعات":
            salary = rng.randint(9000, 24000)
        else:
            salary = rng.randint(10000, 30000)
        days_employed = rng.randint(30, 3650)
        employees.append(
            {
                "Employee Code": f"NC-EMP-DEMO-{index:04d}",
                "Full Name": _synthetic_name(rng),
                "Department": department,
                "Job Title": rng.choice(titles),
                "Monthly Salary EGP": salary,
                "Monthly Incentive EGP": rng.randint(0, 5000),
                "Hire Date": today - timedelta(days=days_employed),
                "Synthetic Demo Data": "YES",
            }
        )
    return employees


def _write_data_sheet(
    workbook: Workbook,
    title: str,
    rows: list[dict[str, Any]],
    table_name: str,
    currency_columns: set[str] | None = None,
) -> None:
    worksheet = workbook.create_sheet(title)
    headers = list(rows[0].keys()) if rows else []
    worksheet.append(headers)
    for row in rows:
        worksheet.append([row.get(header) for header in headers])

    worksheet.freeze_panes = "A2"
    worksheet.sheet_view.showGridLines = False
    worksheet.auto_filter.ref = worksheet.dimensions
    header_fill = PatternFill("solid", fgColor="163B36")
    for cell in worksheet[1]:
        cell.fill = header_fill
        cell.font = Font(color="FFFFFF", bold=True)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    worksheet.row_dimensions[1].height = 32

    for column_index, header in enumerate(headers, start=1):
        samples = [str(header)]
        samples.extend(
            str(worksheet.cell(row=row_index, column=column_index).value or "")
            for row_index in range(2, min(worksheet.max_row, 40) + 1)
        )
        worksheet.column_dimensions[worksheet.cell(row=1, column=column_index).column_letter].width = min(
            max(max(map(len, samples), default=12) + 2, 14), 42
        )
        for row_index in range(2, worksheet.max_row + 1):
            cell = worksheet.cell(row=row_index, column=column_index)
            if isinstance(cell.value, (date, datetime)):
                cell.number_format = "yyyy-mm-dd"
            elif currency_columns and header in currency_columns and isinstance(cell.value, (int, float)):
                cell.number_format = '#,##0.00 "EGP"'
            cell.alignment = Alignment(vertical="top", wrap_text=False)

    if worksheet.max_row > 1 and headers:
        table = Table(displayName=table_name, ref=worksheet.dimensions)
        table.tableStyleInfo = TableStyleInfo(
            name="TableStyleMedium4",
            showFirstColumn=False,
            showLastColumn=False,
            showRowStripes=True,
            showColumnStripes=False,
        )
        worksheet.add_table(table)


def build_workbook(output_path: Path, seed: int = SEED_DEFAULT) -> dict[str, int]:
    rng = random.Random(seed)
    products = generate_products(rng)
    partners = generate_partners(rng)
    sales = generate_sales(rng, products, partners)
    employees = generate_employees(rng)

    workbook = Workbook()
    readme = workbook.active
    readme.title = "README"
    readme.append(["Nova Candles | Enterprise Demo Dataset"])
    readme.append(["هذا الملف يحتوي بيانات اصطناعية للتجربة فقط، وليس بيانات عملاء أو موظفين أو معاملات حقيقية."])
    readme.append(["أرقام الهواتف تبدأ بـ SIM وهي غير صالحة للاتصال."])
    readme.append(["الأوراق: Products (1,000) | Partners (500) | Sales (1,000) | Employees (300)"])
    readme.append(["رفع XML-RPC في هذا السكربت يدعم product.template و res.partner فقط."])
    readme.append(["المبيعات والموظفون والرواتب والمخزون الافتتاحي تبقى في Excel؛ رفعها يتطلب نماذج وصلاحيات وإعدادات إضافية."])
    readme.append(["كل عمليات الكتابة إلى أودو تتطلب الخيار --apply صراحةً؛ بدونه ينفذ السكربت معاينة فقط."])
    readme.column_dimensions["A"].width = 115
    for row in readme.iter_rows():
        row[0].alignment = Alignment(wrap_text=True, vertical="top")
    readme["A1"].font = Font(bold=True, size=16, color="163B36")
    for row_index in range(2, readme.max_row + 1):
        readme.row_dimensions[row_index].height = 30

    _write_data_sheet(
        workbook,
        SHEET_PRODUCTS,
        products,
        "NovaProducts",
        {"Cost EGP", "Wholesale Price EGP", "Retail Price EGP"},
    )
    _write_data_sheet(workbook, SHEET_PARTNERS, partners, "NovaPartners", {"Credit Limit EGP"})
    _write_data_sheet(
        workbook,
        "Sales",
        sales,
        "NovaSales",
        {"Unit Price EGP", "Total EGP"},
    )
    _write_data_sheet(
        workbook,
        "Employees",
        employees,
        "NovaEmployees",
        {"Monthly Salary EGP", "Monthly Incentive EGP"},
    )
    workbook.properties.creator = "Nova Candles - Synthetic ERP Dataset Generator"
    workbook.properties.title = "Nova Candles Enterprise Demo Dataset"
    output_path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(output_path)
    return {
        SHEET_PRODUCTS: len(products),
        SHEET_PARTNERS: len(partners),
        "Sales": len(sales),
        "Employees": len(employees),
    }


def read_records(workbook_path: Path, sheet_name: str) -> list[dict[str, Any]]:
    workbook = load_workbook(workbook_path, read_only=True, data_only=True)
    try:
        candidates = SHEET_ALIASES.get(sheet_name, (sheet_name,))
        actual_sheet = next((name for name in candidates if name in workbook.sheetnames), None)
        if actual_sheet is None:
            available = ", ".join(workbook.sheetnames)
            raise ValueError(f"الورقة {sheet_name} غير موجودة. الأوراق المتاحة: {available}")
        worksheet = workbook[actual_sheet]
        iterator = enumerate(worksheet.iter_rows(values_only=True), start=1)
        name_headers = {
            _normalize_header(alias)
            for aliases in (PARTNER_HEADER_ALIASES["Name"], PRODUCT_HEADER_ALIASES["Name"])
            for alias in aliases
        }
        header_row = next(
            (
                (row_number, values)
                for row_number, values in iterator
                if any(
                    _normalize_header(str(value).strip()) in name_headers
                    for value in values
                    if value is not None
                )
            ),
            None,
        )
        if not header_row:
            return []
        header_row_number, headers = header_row
        normalized_headers = [str(header).strip() if header is not None else "" for header in headers]
        records: list[dict[str, Any]] = []
        # Keep Excel row numbers so an API error can point to the exact source row.
        for row_number, values in iterator:
            if row_number <= header_row_number:
                continue
            if not any(value is not None for value in values):
                continue
            record = dict(zip(normalized_headers, values))
            record["__row__"] = row_number
            records.append(record)
        return records
    finally:
        workbook.close()


def _normalize_header(value: str) -> str:
    return " ".join(value.replace("_", " ").replace("\n", " ").split()).casefold()


def _canonicalize_row(
    source: dict[str, Any],
    aliases: dict[str, tuple[str, ...]],
) -> dict[str, Any]:
    source_by_normalized_header = {
        _normalize_header(header): value
        for header, value in source.items()
        if header != "__row__"
    }
    normalized: dict[str, Any] = {"__row__": source.get("__row__", "?")}
    for canonical_name, candidate_names in aliases.items():
        for candidate in candidate_names:
            key = _normalize_header(candidate)
            if key in source_by_normalized_header:
                normalized[canonical_name] = source_by_normalized_header[key]
                break
    return normalized


def _required_text(row: dict[str, Any], field_name: str) -> str:
    value = row.get(field_name)
    if value is None or not str(value).strip():
        raise ValueError(f"مطلوب حقل {field_name}")
    return str(value).strip()


def _optional_text(row: dict[str, Any], field_name: str) -> str:
    value = row.get(field_name)
    return "" if value is None else str(value).strip()


def _amount(row: dict[str, Any], field_name: str, required: bool = True) -> float | None:
    value = row.get(field_name)
    if value is None or str(value).strip() == "":
        if required:
            raise ValueError(f"مطلوب حقل {field_name}")
        return None
    amount = float(value)
    if not math.isfinite(amount) or amount < 0:
        raise ValueError(f"قيمة غير صالحة في {field_name}: {value}")
    return amount


class OdooRPC:
    """Small authenticated XML-RPC client; credentials are supplied via environment variables."""

    def __init__(self, url: str, database: str, username: str, api_key: str) -> None:
        self.url = url.rstrip("/")
        self.database = database
        self.api_key = api_key
        self.common = xmlrpc.client.ServerProxy(f"{self.url}/xmlrpc/2/common", allow_none=True)
        self.models = xmlrpc.client.ServerProxy(f"{self.url}/xmlrpc/2/object", allow_none=True)
        # Authenticate once and reuse uid for all model operations.
        self.uid = self.common.authenticate(database, username, api_key, {})
        if not self.uid:
            raise RuntimeError("فشل تسجيل الدخول إلى أودو؛ راجع URL وDB واسم المستخدم وAPI key.")

    def call(
        self,
        model: str,
        method: str,
        args: list[Any] | None = None,
        kwargs: dict[str, Any] | None = None,
    ) -> Any:
        return self.models.execute_kw(
            self.database,
            self.uid,
            self.api_key,
            model,
            method,
            args or [],
            kwargs or {},
        )

    def close(self) -> None:
        self.common.close()
        self.models.close()


def _model_fields(client: OdooRPC, model: str) -> dict[str, dict[str, Any]]:
    return client.call(
        model,
        "fields_get",
        [],
        {"attributes": ["type", "string", "readonly", "selection"]},
    )


def _is_writable(fields: dict[str, dict[str, Any]], field_name: str) -> bool:
    return field_name in fields and not fields[field_name].get("readonly", False)


def _selection_has(fields: dict[str, dict[str, Any]], field_name: str, value: str) -> bool:
    if field_name not in fields:
        return False
    options = fields[field_name].get("selection") or []
    return value in {option[0] for option in options}


def _existing_by_key(
    client: OdooRPC,
    model: str,
    key_field: str,
    keys: list[str],
) -> dict[str, int]:
    if not keys:
        return {}
    records = client.call(
        model,
        "search_read",
        [[(key_field, "in", keys)]],
        {"fields": ["id", key_field], "limit": len(keys) + 100},
    )
    return {str(record[key_field]): int(record["id"]) for record in records if record.get(key_field)}


def _ensure_categories(
    client: OdooRPC,
    names: set[str],
    apply: bool,
) -> tuple[dict[str, int], list[str]]:
    if not names:
        return {}, []
    records = client.call(
        "product.category",
        "search_read",
        [[("name", "in", sorted(names))]],
        {"fields": ["id", "name"], "limit": len(names) + 100},
    )
    category_ids = {str(record["name"]): int(record["id"]) for record in records}
    missing = sorted(names - set(category_ids))
    errors: list[str] = []
    if missing and apply:
        for name in missing:
            try:
                category_id = client.call("product.category", "create", [{"name": name}])
                category_ids[name] = int(category_id)
            except Exception as exc:
                message = f"تعذر إنشاء فئة المنتج {name!r}: {exc}"
                print(f"WARNING: {message}", file=sys.stderr)
                errors.append(message)
    elif missing:
        errors.extend(f"الفئة {name!r} ستُنشأ عند التنفيذ الفعلي." for name in missing)
    return category_ids, errors


def _upsert_records(
    client: OdooRPC,
    model: str,
    key_field: str,
    records: list[dict[str, Any]],
    total_rows: int,
    row_errors: list[str],
    apply: bool,
) -> dict[str, int]:
    unique_records: list[dict[str, Any]] = []
    seen_keys: set[str] = set()
    for record in records:
        key = record["key"]
        if key in seen_keys:
            message = _record_error(
                record["sheet"], record["row"], key, "مفتاح مكرر داخل ملف Excel؛ تم تجاوز الصف المكرر."
            )
            row_errors.append(message)
            continue
        seen_keys.add(key)
        unique_records.append(record)

    existing = _existing_by_key(client, model, key_field, [record["key"] for record in unique_records])
    to_create = [record for record in unique_records if record["key"] not in existing]
    to_update = [record for record in unique_records if record["key"] in existing]
    created = 0
    updated = 0
    if not apply:
        print(
            f"معاينة {model}: إجمالي {total_rows}، صالح {len(unique_records)}، "
            f"جديد {len(to_create)}، تحديث {len(to_update)}، فشل تحقق {len(row_errors)}."
        )
        return {"total": total_rows, "created": 0, "updated": 0, "failed": len(row_errors)}

    for record in to_update:
        try:
            client.call(model, "write", [[existing[record["key"]]], record["values"]])
            updated += 1
        except Exception as exc:
            message = _record_error(record["sheet"], record["row"], record["key"], exc)
            row_errors.append(message)

    for start in range(0, len(to_create), BATCH_SIZE):
        batch = to_create[start : start + BATCH_SIZE]
        try:
            values = [record["values"] for record in batch]
            client.call(model, "create", [values])
            created += len(batch)
        except Exception as batch_error:
            # Retry each row independently when one bad row rolls back the batch.
            print(
                f"WARNING: دفعة {model} فشلت ({batch_error})؛ تتم إعادة المحاولة صفًا بصف.",
                file=sys.stderr,
            )
            for record in batch:
                try:
                    current = _existing_by_key(client, model, key_field, [record["key"]])
                    if record["key"] in current:
                        client.call(model, "write", [[current[record["key"]]], record["values"]])
                        updated += 1
                    else:
                        client.call(model, "create", [record["values"]])
                        created += 1
                except Exception as exc:
                    message = _record_error(record["sheet"], record["row"], record["key"], exc)
                    row_errors.append(message)

    failed = len(row_errors)
    print(
        f"التقرير النهائي {model}: إجمالي {total_rows}، نجاح {created + updated} "
        f"(إنشاء {created}، تحديث {updated})، فشل {failed}."
    )
    return {"total": total_rows, "created": created, "updated": updated, "failed": failed}


def _record_error(sheet: str, row_number: Any, key: Any, error: Any) -> str:
    message = f"{sheet}، صف Excel {row_number}، المفتاح {key!r}: {error}"
    print(f"ERROR: {message}", file=sys.stderr)
    return message


def _prepared_record(sheet: str, source: dict[str, Any], key: str, values: dict[str, Any]) -> dict[str, Any]:
    return {"sheet": sheet, "row": source.get("__row__", "?"), "key": key, "values": values}


def upload_products(client: OdooRPC, rows: list[dict[str, Any]], apply: bool) -> dict[str, int]:
    model = "product.template"
    fields = _model_fields(client, model)
    if not _is_writable(fields, "default_code"):
        raise RuntimeError("الحقل default_code غير قابل للكتابة؛ لا يمكن ضمان upsert آمن للمنتجات.")

    prepared: list[dict[str, Any]] = []
    row_errors: list[str] = []
    category_names: set[str] = set()
    skipped: set[str] = set()
    has_opening_stock = False
    for index, source in enumerate(rows, start=2):
        row = _canonicalize_row(source, PRODUCT_HEADER_ALIASES)
        try:
            sku = _required_text(row, "SKU")
            name = _required_text(row, "Name")
            retail_price = _amount(row, "Retail Price EGP")
            cost = _amount(row, "Cost EGP")
            stock = _amount(row, "Quantity On Hand", required=False)
            has_opening_stock = has_opening_stock or bool(stock)
            category_name = _optional_text(row, "Product Category")
            if category_name:
                category_names.add(category_name)

            values: dict[str, Any] = {}
            candidates = {
                "name": name,
                "default_code": sku,
                "list_price": retail_price,
                "standard_price": cost,
            }
            description = " ".join(
                part for part in (
                    _optional_text(row, "Sales Description"),
                    f"الخامات: {_optional_text(row, 'Materials')}" if _optional_text(row, "Materials") else "",
                    f"الزيت العطري: {_optional_text(row, 'Fragrance Oil')}" if _optional_text(row, "Fragrance Oil") else "",
                ) if part
            )
            if description:
                candidates["description_sale"] = description
            for field_name, value in candidates.items():
                if _is_writable(fields, field_name):
                    values[field_name] = value
                else:
                    skipped.add(field_name)

            product_type = _optional_text(row, "Product Type").casefold()
            if _is_writable(fields, "is_storable"):
                values["is_storable"] = True
            for type_field in ("type", "detailed_type"):
                if _is_writable(fields, type_field) and _selection_has(fields, type_field, "consu"):
                    values[type_field] = "consu"
                    break

            wick_field = "candle_wick_type"
            wick_code = _optional_text(row, "Wick Type Code").casefold()
            wick_label = _optional_text(row, "Wick Type").casefold()
            if not wick_code:
                if "خشب" in wick_label or "wood" in wick_label:
                    wick_code = "wood"
                elif "قطن" in wick_label or "cotton" in wick_label:
                    wick_code = "cotton"
            if _is_writable(fields, wick_field) and _selection_has(fields, wick_field, wick_code):
                values[wick_field] = wick_code
            if product_type and not (_is_writable(fields, "is_storable") or "type" in fields or "detailed_type" in fields):
                skipped.add("Product Type (no compatible Odoo field)")

            wholesale = _amount(row, "Wholesale Price EGP", required=False)
            wholesale_field = next(
                (field_name for field_name in ("wholesale_price", "x_wholesale_price") if _is_writable(fields, field_name)),
                None,
            )
            if wholesale is not None and wholesale_field:
                values[wholesale_field] = wholesale
            elif wholesale is not None:
                skipped.add("Wholesale Price EGP (no compatible Odoo field)")

            product_record = _prepared_record(SHEET_PRODUCTS, source, sku, values)
            product_record["category"] = category_name
            prepared.append(product_record)
        except (TypeError, ValueError, OverflowError) as exc:
            row_errors.append(_record_error(SHEET_PRODUCTS, source.get("__row__", index), source.get("SKU", ""), exc))

    categories, category_warnings = _ensure_categories(client, category_names, apply)
    if category_warnings:
        for warning in category_warnings:
            print(f"WARNING: {warning}", file=sys.stderr)
    if _is_writable(fields, "categ_id"):
        for record in prepared:
            category_name = record["category"]
            if category_name in categories:
                record["values"]["categ_id"] = categories[category_name]
            elif category_name:
                skipped.add("Product Category (unresolved)")

    if has_opening_stock:
        skipped.add("Quantity On Hand (opening stock remains in Excel; use Odoo inventory adjustments)")
    if skipped:
        print("حقول لم تُرسل أو بقيت في Excel: " + ", ".join(sorted(skipped)))
    return _upsert_records(client, model, "default_code", prepared, len(rows), row_errors, apply)


def upload_partners(client: OdooRPC, rows: list[dict[str, Any]], apply: bool) -> dict[str, int]:
    model = "res.partner"
    fields = _model_fields(client, model)
    if not _is_writable(fields, "ref"):
        raise RuntimeError("الحقل ref غير قابل للكتابة؛ لا يمكن ضمان upsert آمن للعملاء.")

    prepared: list[dict[str, Any]] = []
    row_errors: list[str] = []
    skipped: set[str] = set()
    for index, source in enumerate(rows, start=2):
        row = _canonicalize_row(source, PARTNER_HEADER_ALIASES)
        try:
            name = _required_text(row, "Name")
            city = _optional_text(row, "City")
            if not city:
                district = _optional_text(row, "District")
                governorate = _optional_text(row, "Governorate")
                city = ", ".join(value for value in (district, governorate) if value)
            phone = _optional_text(row, "Phone")
            segment = _optional_text(row, "Customer Segment")
            customer_code = _optional_text(row, "Customer Code")
            if not customer_code:
                stable_value = "|".join((name.casefold(), city.casefold(), phone.casefold()))
                customer_code = "AUTO-" + hashlib.sha256(stable_value.encode("utf-8")).hexdigest()[:16].upper()

            location_parts = [part for part in (city, "مصر" if city else "") if part]
            note_parts = []
            if segment:
                note_parts.append(f"التصنيف: {segment}")
            channel = _optional_text(row, "Sales Channel")
            if channel:
                note_parts.append(f"قناة البيع: {channel}")
            contact = _optional_text(row, "Contact Person")
            if contact:
                note_parts.append(f"مسؤول التواصل: {contact}")

            values: dict[str, Any] = {"name": name, "ref": customer_code}
            if _is_writable(fields, "customer_rank"):
                values["customer_rank"] = 1
            if phone:
                values["phone"] = phone
            if location_parts:
                values["city"] = ", ".join(location_parts)
            if note_parts:
                values["comment"] = " | ".join(note_parts)
            values = {field_name: value for field_name, value in values.items() if _is_writable(fields, field_name)}

            is_individual = any(marker in segment.casefold() for marker in ("فرد", "individual"))
            company_type = "person" if is_individual else "company"
            if _is_writable(fields, "company_type") and _selection_has(fields, "company_type", company_type):
                values["company_type"] = company_type
            elif _is_writable(fields, "is_company"):
                values["is_company"] = not is_individual

            credit_limit = _amount(row, "Credit Limit EGP", required=False)
            credit_field = next(
                (field_name for field_name in ("credit_limit", "x_credit_limit") if _is_writable(fields, field_name)),
                None,
            )
            if credit_limit is not None and credit_field:
                values[credit_field] = credit_limit
            elif credit_limit is not None:
                skipped.add("Credit Limit EGP (no compatible Odoo field)")
            prepared.append(_prepared_record(SHEET_PARTNERS, source, customer_code, values))
        except (TypeError, ValueError, OverflowError) as exc:
            row_errors.append(
                _record_error(SHEET_PARTNERS, source.get("__row__", index), source.get("Customer Code", ""), exc)
            )

    if skipped:
        print("حقول لم تُرسل أو بقيت في Excel: " + ", ".join(sorted(skipped)))
    return _upsert_records(client, model, "ref", prepared, len(rows), row_errors, apply)


def _parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(
        description="إنشاء بيانات Nova Candles اصطناعية في Excel، مع رفع اختياري للمنتجات والعملاء إلى أودو."
    )
    parser.add_argument("--output", type=Path, default=Path("Nova_Candles_Enterprise_Demo.xlsx"))
    parser.add_argument("--input-workbook", type=Path, help="استخدم ملف Excel موجودًا بدل إنشاء ملف جديد.")
    parser.add_argument("--seed", type=int, default=SEED_DEFAULT)
    parser.add_argument("--upload-products", action="store_true", help="معاينة/رفع ورقة Products إلى product.template.")
    parser.add_argument("--upload-partners", action="store_true", help="معاينة/رفع ورقة Partners إلى res.partner.")
    parser.add_argument("--apply", action="store_true", help="نفّذ الكتابة الفعلية؛ بدونه تكون العملية معاينة فقط.")
    return parser.parse_args()


def main() -> int:
    args = _parse_args()
    upload_requested = args.upload_products or args.upload_partners
    if args.apply and not upload_requested:
        print("--apply يحتاج --upload-products أو --upload-partners.", file=sys.stderr)
        return 2

    workbook_path = args.input_workbook
    if workbook_path:
        if not workbook_path.is_file():
            print(f"ملف Excel غير موجود: {workbook_path}", file=sys.stderr)
            return 2
        print(f"سيتم استخدام الملف: {workbook_path}")
    else:
        counts = build_workbook(args.output, args.seed)
        workbook_path = args.output
        print(f"تم إنشاء الملف: {workbook_path.resolve()}")
        print("أعداد السجلات: " + ", ".join(f"{name}={count:,}" for name, count in counts.items()))

    if not upload_requested:
        return 0

    # Upload is opt-in; without --apply this path only authenticates and previews.
    url = os.getenv("ODOO_URL", "").strip()
    database = os.getenv("ODOO_DB", "").strip()
    username = os.getenv("ODOO_USERNAME", "").strip()
    api_key = os.getenv("ODOO_API_KEY", "").strip() or os.getenv("ODOO_PASSWORD", "").strip()
    missing = [name for name, value in {
        "ODOO_URL": url,
        "ODOO_DB": database,
        "ODOO_USERNAME": username,
        "ODOO_API_KEY أو ODOO_PASSWORD": api_key,
    }.items() if not value]
    if missing:
        print("متغيرات البيئة المطلوبة غير محددة: " + ", ".join(missing), file=sys.stderr)
        return 2

    client: OdooRPC | None = None
    try:
        client = OdooRPC(url, database, username, api_key)
        mode = "تنفيذ فعلي" if args.apply else "معاينة فقط (لم يتم تعديل أودو)"
        print(f"اتصال أودو ناجح. الوضع: {mode}.")
        if args.upload_products:
            products = read_records(workbook_path, SHEET_PRODUCTS)
            upload_products(client, products, args.apply)
        if args.upload_partners:
            partners = read_records(workbook_path, SHEET_PARTNERS)
            upload_partners(client, partners, args.apply)
        if not args.apply:
            print("للتنفيذ الفعلي، أعد الأمر مع --apply بعد مراجعة المعاينة.")
        return 0
    except (OSError, ValueError, RuntimeError, xmlrpc.client.Error) as exc:
        print(f"فشل التنفيذ: {exc}", file=sys.stderr)
        return 1
    finally:
        if client is not None:
            client.close()


if __name__ == "__main__":
    raise SystemExit(main())